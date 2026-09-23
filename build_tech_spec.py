from __future__ import annotations

import csv
import re
import shutil
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Tech Spec"
DS_DIR = OUT / "Data sheets"


@dataclass
class Part:
    ident: str
    category: str
    item: str
    qty: str
    maturity: str
    candidate: str
    purpose: str
    requirements: str
    interfaces: str
    acceptance: str
    dependencies: str
    source: str


def p(category, item, qty, maturity, candidate, purpose, requirements,
      interfaces, acceptance, dependencies="Nenhuma adicional.", source="Ficha do fornecedor a exigir antes da compra."):
    ident = f"DS-{len(PARTS)+1:03d}"
    PARTS.append(Part(ident, category, item, qty, maturity, candidate, purpose,
                      requirements, interfaces, acceptance, dependencies, source))


PARTS: list[Part] = []

# A — Aviónica e missão
p("Aviónica", "Controlador de voo", "1", "Arquitetura aceita", "Holybro Pixhawk 6X",
  "Controle de voo, navegação, failsafes e registro primário.",
  "Autopiloto classe H7, IMUs redundantes, barômetros redundantes, CAN, Ethernet e entradas de alimentação redundantes.",
  "MAVLink/CAN/UART; 5 V regulados; servo rail isolado.",
  "Teste de alimentação A/B, sensores, logging, watchdog, RTL e perda de enlace.",
  "Firmware ArduPilot/PX4 e integração do sistema de energia.",
  "https://docs.holybro.com/autopilot/pixhawk-6x/technical-specification")
p("Aviónica", "Computador de missão", "1", "Arquitetura aceita", "Raspberry Pi 5 — 8 GB",
  "Visão computacional, comunicações de alto nível e processamento de missão.",
  "BCM2712, 8 GB, alimentação 5 V/5 A, refrigeração ativa e desligamento controlado.",
  "Ethernet/USB/UART com Pixhawk; câmera CSI; armazenamento local.",
  "Burn-in de 4 h, carga de CPU, perda/retorno de energia e temperatura no compartimento.",
  "Conversores DC/DC redundantes e solução térmica.",
  "https://datasheets.raspberrypi.com/rpi5/raspberry-pi-5-product-brief.pdf")
p("Aviónica", "Refrigerador do computador", "1", "Candidato", "Raspberry Pi Active Cooler",
  "Manter o computador de missão abaixo do limite térmico.",
  "Compatível com Raspberry Pi 5; operação contínua na temperatura interna prevista.",
  "Conector de fan do Raspberry Pi 5; fluxo de ar desobstruído.",
  "Teste térmico em solo com CPU a 100% e fuselagem fechada.",
  "Projeto de dutos/ventilação.", "https://www.raspberrypi.com/products/active-cooler/")
p("Aviónica", "Módulo GNSS/RTK + magnetômetro", "1", "Candidato", "Holybro H-RTK NEO-F9P + RM3100",
  "Posição, velocidade, rumo e possibilidade de RTK.",
  "Dupla banda GNSS, RTK, magnetômetro RM3100; 4,75–5,25 V; massa aproximada 39 g.",
  "CAN/UART conforme variante; montagem afastada de potência e ignição.",
  "Aquisição a céu aberto, comparação de rumo, interferência com motor a gasolina ligado.",
  "Mastro e plano de aterramento eletromagnético.",
  "https://docs.holybro.com/gps-and-rtk-system/h-rtk-neo-f9p-series-rm3100-compass/specification")
p("Aviónica", "Sensor de velocidade do ar + pitot", "1", "Candidato", "Holybro MS5525",
  "Medição de velocidade indicada e proteção do envelope.",
  "Sensor diferencial digital, faixa adequada ao voo de 15–40 m/s, tubos resistentes a combustível/calor.",
  "I²C/CAN via interface; tomadas de pressão total/estática.",
  "Calibração zero, teste de vazamento, comparação em túnel/veículo.",
  "Posição do pitot fora da esteira das hélices.",
  "https://docs.holybro.com/peripherals/ms4525do-and-ms5525dso-i2c-airspeed-sensor")
p("Aviónica", "Altímetro laser", "1", "Candidato", "LightWare SF11/C",
  "Altura precisa na decolagem e no pouso VTOL.",
  "Faixa 0,2–100 m, 20 Hz, precisão aproximada ±0,1 m, 5 V, massa aproximada 35 g.",
  "Serial/I²C; campo de visão desobstruído para baixo.",
  "Teste em superfícies clara/escura, poeira e inclinação; comparação com referência.",
  "Janela óptica e suporte antivibração.",
  "https://lightwarelidar.com/wp-content/uploads/2025/07/SF11-Laser-Altimeter-Manual-Rev-10.pdf")
p("Missão", "Câmera RGB", "1", "Candidato", "Raspberry Pi Camera Module 3 Wide",
  "Busca visual, registro e apoio à navegação.",
  "Sony IMX708, 12 MP, autofoco, variante wide; proteção óptica substituível.",
  "CSI para Raspberry Pi 5.",
  "Foco, exposição, vibração, rolling shutter e gravação sustentada.",
  "Cabo CSI com comprimento final e janela frontal.",
  "https://datasheets.raspberrypi.com/camera/camera-module-3-product-brief.pdf")
p("Missão", "Câmera térmica", "1", "Opcional / candidato", "Teledyne FLIR Boson 320 ou 640",
  "Detecção térmica de pessoas em missão de resgate.",
  "Núcleo 12 µm, resolução selecionada por alcance/ orçamento, consumo e lente compatíveis com a missão.",
  "USB/serial/vídeo conforme variante; montagem estabilizada ou fixa.",
  "Teste de alcance de detecção, latência, temperatura e gravação sincronizada.",
  "Definir requisito operacional de alcance e lente.", "https://www.oem.flir.com/products/boson/")
p("Aviónica", "Armazenamento de missão", "2", "TBD", "microSD High Endurance ou NVMe industrial",
  "Sistema operacional e logs redundantes.",
  "Alta resistência de escrita, temperatura industrial preferencial, capacidade ≥128 GB.",
  "microSD ou M.2 HAT; imagem de recuperação documentada.",
  "Teste SMART/velocidade, 8 h de escrita e recuperação após corte de energia.",
  "Escolha final do gabinete do Raspberry Pi.")
p("Aviónica", "Buzzer e localizador", "1", "TBD", "Buzzer ativo + beacon independente",
  "Localização da aeronave e indicação de estado.",
  "Audível a longa distância; alimentação independente ou reserva; baixo consumo em espera.",
  "GPIO/PWM e bateria reserva.",
  "Teste de autonomia, alcance audível e acionamento após falha do barramento.")
p("Aviónica", "Módulo de relógio RTC", "1", "TBD", "RTC industrial com bateria/supercapacitor",
  "Preservar tempo de logs sem GNSS/rede.",
  "Baixa deriva, interface compatível com computador de missão.",
  "I²C; bateria substituível ou supercapacitor.",
  "Teste de deriva por 7 dias e retomada após corte total.")

# B — Comunicação
p("Comunicação", "Rádio telemetria de bordo", "1", "Candidato", "RFD900x-BR",
  "Telemetria MAVLink de longo alcance.",
  "Versão homologada/adequada ao Brasil, 5 V, diversidade e potência configurável.",
  "UART para Pixhawk; duas antenas afastadas de fontes EMI.",
  "Teste de enlace, failsafe, corrente em TX máximo e alcance progressivo.",
  "Confirmar homologação ANATEL e plano de frequências.",
  "https://rfdesign.com.au/wp-content/uploads/2024/04/RFD900x-DataSheet-V1.2.pdf")
p("Comunicação", "Rádio telemetria de solo", "1", "Candidato", "RFD900x-BR USB",
  "Par do rádio embarcado para a estação de controle.",
  "Mesma banda/firmware; interface USB; antenas compatíveis.",
  "USB para estação de solo.",
  "Teste de pareamento, atualização, RSSI e recuperação de perda.",
  "Mesmo lote/região do rádio de bordo.",
  "https://store.rfdesign.com.au/all-products/")
p("Comunicação", "Conjunto de antenas e cabos RF", "1 kit", "TBD", "Antenas 900 MHz + cabos coaxiais",
  "Radiar/receber telemetria com baixa perda.",
  "Impedância 50 Ω, VSWR documentado, conectores corretos, cabos de baixa perda.",
  "SMA/RP-SMA conforme modem; montagem ortogonal/diversidade.",
  "Medição de continuidade/VSWR e teste de alcance.",
  "Definir instalação e conectores antes da compra.")
p("Comunicação", "Modem 4G", "1", "Candidato", "Waveshare SIM7600G-H 4G HAT",
  "Canal secundário de dados e supervisão quando houver cobertura.",
  "LTE Cat-4 global; confirmar bandas da operadora brasileira; GNSS do modem não substitui GNSS de voo.",
  "USB para Raspberry Pi; alimentação 5 V com picos suportados.",
  "Teste de bandas, handover, consumo de pico e perda de rede.",
  "SIM/plano de dados e cobertura operacional.", "https://www.waveshare.com/SIM7600G-H-4G-HAT.htm")
p("Comunicação", "Antenas LTE/GNSS do modem", "1 kit", "TBD", "Antenas compatíveis com SIM7600",
  "Comunicação celular e GNSS auxiliar.",
  "Bandas brasileiras, 50 Ω, conectores corretos e instalação distante da ignição.",
  "U.FL/SMA conforme chicote.",
  "VSWR, cobertura e interferência com demais rádios.")

# C — Propulsão VTOL e tilt
p("Propulsão VTOL", "Motor elétrico de sustentação", "4", "Candidato — compra bloqueada", "T-Motor U10 II KV100",
  "Sustentação VTOL e, quando comandado, contribuição elétrica na transição.",
  "12S; massa aproximada 415 g; compatível com hélice 30×10,5; margem térmica e empuxo final validados em bancada.",
  "ESC trifásico; flange estrutural; sensor de rotação recomendado.",
  "Curva empuxo/corrente, temperatura, vibração, balanceamento e falha de um motor.",
  "MTOW fechado e razão empuxo/peso VTOL ≥1,5.", "https://store.tmotor.com/product/u10-2-u-efficiency.html")
p("Propulsão VTOL", "ESC dos motores de sustentação", "4", "Candidato — compra bloqueada", "T-Motor ALPHA 80A 12S",
  "Comutação e controle dos quatro motores VTOL.",
  "12S, corrente contínua e refrigeração compatíveis com curva real; telemetria desejável.",
  "Bateria 12S, motor U10 II, PWM/DShot conforme integração.",
  "Bancada em potência máxima, rampa, corte, perda de sinal e temperatura.",
  "Curva final de hélice/motor e projeto de ventilação.", "https://store.tmotor.com/product/alpha-80a-12s-esc.html")
p("Propulsão VTOL", "Hélice VTOL CW 30×10,5", "2 + 1 reserva", "Candidato — compra bloqueada", "T-Motor G30×10.5 CW",
  "Gerar empuxo vertical com sentido de rotação balanceado.",
  "Carbono, diâmetro/passo compatíveis, limite de RPM documentado.",
  "Hub do U10 II; envelope livre em todas as posições de tilt.",
  "Balanceamento estático/dinâmico, inspeção e prova de overspeed controlada.",
  "Validação geométrica de folga.")
p("Propulsão VTOL", "Hélice VTOL CCW 30×10,5", "2 + 1 reserva", "Candidato — compra bloqueada", "T-Motor G30×10.5 CCW",
  "Gerar empuxo vertical com cancelamento de torque.",
  "Carbono, geometria espelhada e limite de RPM documentado.",
  "Hub do U10 II; envelope livre em todas as posições de tilt.",
  "Balanceamento estático/dinâmico, inspeção e prova de overspeed controlada.",
  "Validação geométrica de folga.")
p("Tilt-rotor", "Atuador de inclinação da nacele", "4", "TBD crítico", "Atuador eletromecânico industrial com feedback absoluto",
  "Girar motores dianteiros para frente e traseiros para trás somente com hélices seguras.",
  "Torque com fator ≥2 sobre pior carga aerodinâmica/inercial; feedback absoluto; retenção sem energia; baixa folga.",
  "Alimentação dedicada; CAN/PWM; eixo e mancais da nacele.",
  "Ciclagem ≥10.000, torque, folga, vibração, falha de energia e posição.",
  "Cargas de transição e projeto mecânico dos pivôs.")
p("Tilt-rotor", "Trava mecânica da nacele", "4", "TBD crítico", "Pino/trava fail-safe com sensor duplo",
  "Imobilizar cada nacele em VTOL e cruzeiro sem depender do torque do atuador.",
  "Fail-safe, confirmação independente de engate, resistência com fator estrutural definido.",
  "Interlock do controlador; batentes estruturais.",
  "Carga limite/última, engate contaminado, falha simples de sensor e 10.000 ciclos.",
  "Geometria do pivô e cargas finais.")
p("Tilt-rotor", "Sensor de indexação/parada da hélice", "4", "TBD crítico", "Encoder/Hall redundante",
  "Confirmar RPM zero e posição angular segura antes do movimento da nacele.",
  "Detecção redundante de velocidade zero e ângulo; imunidade EMI; resolução suficiente para folga estrutural.",
  "Motor/ESC e controlador de intertravamento.",
  "Falha de sensor, falso zero, parada fora da janela e vibração.",
  "Definir posição segura de cada pá no CAD.")
p("Tilt-rotor", "Mancais e eixo do pivô", "4 kits", "TBD crítico", "Rolamentos selados + eixo 7075/aço",
  "Suportar a nacele e transferir empuxo/cargas para a estrutura.",
  "Vida e carga radial/axial calculadas; baixa folga; vedação contra poeira/água.",
  "Nacele, trava, atuador e longarina/boom.",
  "Carga limite, inspeção de folga e ciclagem.",
  "Análise estrutural e desenhos liberados.")

# D — Cruzeiro a gasolina e geração
p("Propulsão cruzeiro", "Motor a gasolina pusher", "1", "Candidato — compra bloqueada", "DLE35RA",
  "Propulsão principal de cruzeiro após atingir altitude/condições de transição.",
  "34,9 cm³; potência declarada 4,1 hp a 8.500 rpm; ignição 4,8–8,4 V; combustível 30:1 conforme fabricante.",
  "Firewall traseiro, hélice pusher, tanque, servo de aceleração, ignição e escapamento.",
  "Amaciamento, curva de empuxo/consumo, partida em voo simulada, temperatura e vibração.",
  "Arrasto/potência de cruzeiro, partida pelo vento e orçamento de combustível.", "https://dlengine.com/en/rcengine/dle35ra")
p("Propulsão cruzeiro", "Hélice pusher", "2", "TBD crítico", "19×8 ou 20×8 pusher compatível",
  "Converter potência do DLE35RA em empuxo de cruzeiro.",
  "Orientação pusher real, material e RPM compatíveis; folga de fuselagem/cauda.",
  "Cubo do motor e spinner se aplicável.",
  "Balanceamento, empuxo/consumo, ruído e overspeed; uma unidade como reserva.",
  "Curva motor-hélice e instalação final.")
p("Propulsão cruzeiro", "Servo de aceleração", "1", "TBD", "Servo HV metal gear",
  "Controle do carburador.",
  "Torque e velocidade suficientes, engrenagens metálicas, rolamentos, faixa térmica e vibração.",
  "PWM do controlador; linkagem com mola de retorno para marcha lenta/corte.",
  "Ciclagem, retorno seguro, vibração e perda de sinal.")
p("Propulsão cruzeiro", "Módulo de corte da ignição", "1", "TBD crítico", "Kill switch optoisolado",
  "Desligar o motor por comando ou failsafe.",
  "Isolação elétrica, estado seguro sem sinal, compatível com tensão/corrente da ignição.",
  "Pixhawk/relé de segurança e módulo de ignição.",
  "Corte por RC, autonomia, geofence, perda de alimentação e EMI.")
p("Propulsão cruzeiro", "Sensor de RPM do motor térmico", "1", "TBD", "Sensor Hall/óptico isolado",
  "Confirmar rotação, partida e falhas do motor.",
  "Faixa ≥12.000 rpm, imunidade EMI e saída compatível com autopiloto/computador.",
  "Eixo/ímã ou ignição; entrada isolada.",
  "Comparação com tacômetro e teste com ignição ativa.")
p("Propulsão cruzeiro", "Motor elétrico starter-generator", "1", "TBD crítico", "BLDC dimensionado para partida e geração",
  "Dar partida no motor a gasolina e gerar energia em cruzeiro.",
  "Torque de partida, rotação, potência contínua de geração, refrigeração e massa definidos por ensaio.",
  "Acoplamento unidirecional, controlador bidirecional e barramento elétrico.",
  "Partida a frio/quente, geração contínua, overspeed e falha desacoplada.",
  "Torque de partida do DLE35RA e balanço de energia.")
p("Propulsão cruzeiro", "Controlador bidirecional starter-generator", "1", "TBD crítico", "Inversor/retificador BLDC bidirecional",
  "Controlar partida e regular carga do barramento/bateria.",
  "Tensão 12S, corrente de partida, regeneração controlada, proteção contra sobretensão.",
  "Starter-generator, barramento, BMS/monitor e controlador.",
  "Transições motor/gerador, carga rejeitada, sobretensão e perda de comunicação.",
  "Motor starter-generator e arquitetura de bateria.")
p("Propulsão cruzeiro", "Embreagem/roda livre e transmissão", "1 kit", "TBD crítico", "Roda livre + correia/engrenagem",
  "Acoplar partida/geração sem impor risco mecânico ao motor térmico.",
  "Torque, rotação, vida e contenção dimensionados; guarda física contra fragmentos.",
  "Virabrequim/eixo e starter-generator.",
  "Torque limite, overspeed, desalinhamento, perda de lubrificação e contenção.",
  "Layout mecânico e análise torsional.")
p("Combustível", "Tanque ou bladder de combustível", "1", "TBD crítico", "Capacidade geométrica inicial 4 L",
  "Armazenar gasolina/óleo para a missão.",
  "Compatível com gasolina, respiro seguro, anti-vazamento, sump; capacidade final pelo ensaio de consumo + reserva.",
  "Linhas, filtro, clunk, respiro e sensor de nível.",
  "Estanqueidade, vibração, expansão térmica, alimentação em atitudes e crashworthiness básica.",
  "Consumo medido e requisito de 400 km.")
p("Combustível", "Kit de linhas, filtro, clunk e respiro", "1 kit", "TBD", "Componentes compatíveis com gasolina",
  "Conduzir e filtrar combustível em todas as atitudes.",
  "Tubos resistentes a gasolina, filtro inspecionável, clunk e respiro anti-sifão.",
  "Tanque e carburador.",
  "Vazão, bolhas, vazamento, atitude e vibração.")
p("Combustível", "Válvula de corte de combustível", "1", "TBD crítico", "Válvula fail-safe compatível com gasolina",
  "Isolar o combustível em emergência/manutenção.",
  "Baixa perda de carga, compatibilidade química, posição segura definida.",
  "Linha entre tanque e carburador; comando de segurança.",
  "Vazamento, vazão, falha de energia e acionamento remoto.")
p("Combustível", "Sensor de nível/fluxo de combustível", "1", "TBD", "Sensor compatível com gasolina",
  "Estimar combustível restante e consumo real.",
  "Baixa restrição, precisão calibrável e compatibilidade química.",
  "Linha/tanque; entrada do computador/autopiloto.",
  "Calibração volumétrica, vazão pulsante e temperatura.")

# E — Energia elétrica
p("Energia", "Célula Li-ion 21700", "48", "Candidato — compra bloqueada", "Molicel INR-21700-P45B; pack conceitual 12S4P",
  "Armazenar energia para VTOL, partida e aviónica.",
  "3,6 V, 4,5 Ah e 45 A contínuos por célula declarados; células do mesmo lote, rastreadas e casadas.",
  "Soldagem/união industrial, monitor por célula, contenção e ventilação.",
  "Capacidade, resistência interna, casamento, abuso térmico por análise e ensaio do pack.",
  "Perfil VTOL medido, massa final e projeto profissional do pack.", "https://www.molicel.com/product/inr-21700-p45b/")
p("Energia", "BMS/monitor de células 12S", "1", "TBD crítico", "BMS aeronáutico ou monitor independente",
  "Monitorar tensão/temperatura e proteger carga; estratégia de corte em voo deve ser segura.",
  "12S, sensores térmicos múltiplos, logging, balanceamento; não cortar propulsão abruptamente por falha simples.",
  "Pack, carregador e controlador de energia.",
  "OV/UV/OT, falha de sensor, comunicação e comportamento em voo.",
  "Análise de segurança da arquitetura.")
p("Energia", "Contator principal HV", "1", "TBD crítico", "Contator DC ≥60 V",
  "Isolar a bateria de tração.",
  "Corrente contínua/pico dimensionada, bobina eficiente, contato auxiliar, supressão de arco.",
  "Pack, pré-carga e barramento.",
  "Queda de tensão, aquecimento, abertura sob carga controlada e vida.")
p("Energia", "Circuito de pré-carga", "1", "TBD crítico", "Resistor + relé/contator de pré-carga",
  "Limitar corrente de carga dos capacitores dos ESCs.",
  "Energia do resistor e tempo de pré-carga calculados; confirmação de tensão antes do contator.",
  "Pack, ESCs e contator.",
  "Tempo de carga, falha aberta/fechada e temperatura.")
p("Energia", "Fusível principal HV", "1 + 2 reservas", "TBD crítico", "Fusível DC aeronáutico/EV",
  "Proteção contra curto no barramento principal.",
  "Tensão DC e capacidade de interrupção adequadas; curva I²t coordenada com cabos/contator.",
  "Próximo ao pack, suporte protegido.",
  "Verificação de seletividade e inspeção do suporte.")
p("Energia", "Distribuidor/barramento de potência HV", "1", "Sob medida", "PDB/busbar 12S",
  "Distribuir energia aos quatro ESCs e starter-generator.",
  "Corrente e temperatura dimensionadas, proteção física e pontos de teste.",
  "Pack, fusíveis de ramo, ESCs e conversores.",
  "Queda de tensão, termografia e curto controlado por análise.",
  "Correntes medidas de todos os ramos.")
p("Energia", "Módulo de medição de potência", "1", "Candidato", "Holybro PM08-CAN ou solução equivalente",
  "Medir tensão/corrente e alimentar o controlador de voo.",
  "Faixa de tensão 12S, corrente compatível, CAN; saída 5,2 V não deve alimentar sozinha o Raspberry Pi.",
  "CAN e Power do Pixhawk; shunt no barramento.",
  "Calibração de V/I, sobrecorrente e redundância de alimentação.",
  "Definir corrente máxima do barramento.", "https://docs.holybro.com/power-module-and-pdb/power-module-comparison")
p("Energia", "Conversor DC/DC 5 V 5 A isolado", "2", "TBD crítico", "Conversores industriais redundantes",
  "Alimentar computador de missão e periféricos sem ruído da propulsão.",
  "Entrada 12S completa, 5,1 V/5 A contínuos cada, isolamento/EMI, compartilhamento ou comutação segura.",
  "Barramento 12S, Raspberry Pi e periféricos.",
  "Carga 125%, ripple, transientes, temperatura e falha de um canal.")
p("Energia", "Conversor DC/DC 12 V", "1", "TBD", "Conversor industrial 12S→12 V",
  "Alimentar acessórios de 12 V.",
  "Potência conforme cargas, proteção e EMI.",
  "Barramento 12S e acessórios.",
  "Carga, ripple, temperatura e curto.")
p("Energia", "Bateria reserva de aviónica", "1", "TBD crítico", "Pack LiFePO4 independente",
  "Manter controle, navegação e telemetria após falha do barramento principal.",
  "Energia para pouso/recuperação segura, química estável, monitoramento e isolação.",
  "Power 2 do Pixhawk e cargas essenciais via OR-ing.",
  "Autonomia, transferência sem interrupção e alerta de baixa tensão.",
  "Definir duração de emergência.")
p("Energia", "Módulo ideal-diode/OR-ing", "1", "TBD crítico", "OR-ing de duas fontes",
  "Selecionar alimentação primária/reserva sem retorno de corrente.",
  "Tensão/corrente compatíveis, baixa queda e telemetria de estado.",
  "Fontes principal/reserva e aviónica.",
  "Falha de cada fonte, curto e transição.")
p("Energia", "Conectores antifaísca HV", "1 kit", "TBD", "AS150 ou equivalente qualificado",
  "Conexão segura do pack e manutenção.",
  "Tensão/corrente, polarização, proteção contra toque e ciclos suficientes.",
  "Pack/barramento; identificação mecânica.",
  "Aquecimento, ciclos, retenção e inspeção.")
p("Energia", "Cabos de potência e sinal", "1 lote", "TBD", "Cabos flexíveis aeronáuticos",
  "Interligar energia, sensores e comunicação.",
  "Bitolas calculadas, isolamento térmico/químico, pares trançados/blindados onde necessário.",
  "Terminais crimpados, separação potência/sinal.",
  "Pull test, continuidade, isolamento, queda de tensão e inspeção 100%.",
  "Diagrama elétrico e roteamento finais.")
p("Energia", "Sensores de temperatura", "8", "TBD", "NTC/PT1000 digitais/analógicos",
  "Monitorar pack, ESCs, DC/DC e compartimento do motor.",
  "Faixa e precisão compatíveis; fixação segura; canais suficientes.",
  "BMS/computador/autopiloto.",
  "Calibração em dois pontos e teste de falha aberta/curta.")

# F — Atuadores de controle
p("Controles", "Servo de aileron HV", "2 + 1 reserva", "TBD crítico", "Servo digital industrial metal gear",
  "Acionar ailerons esquerdo e direito.",
  "Torque calculado com fator ≥1,5, engrenagens metálicas, rolamentos, baixa folga e faixa térmica.",
  "PWM/CAN; 6–8,4 V; pushrod/hinge.",
  "Torque, velocidade, centragem, folga e 20.000 ciclos sob carga.",
  "Cálculo de momento de dobradiça.")
p("Controles", "Servo de ruddervator da cauda V", "2 + 1 reserva", "TBD crítico", "Servo digital industrial metal gear",
  "Controlar arfagem e guinada na empenagem em V.",
  "Torque calculado, baixa folga e redundância funcional por superfícies separadas.",
  "Mixer do autopiloto; 6–8,4 V; linkagens.",
  "Deflexão, torque, centragem e falha de um servo.",
  "Cálculo de cargas da cauda V a 40°.")
p("Controles", "Servo de porta/liberação de carga", "1", "TBD", "Servo ou atuador com trava",
  "Abrir porta ou liberar carga somente sob comando autorizado.",
  "Trava mecânica, indicação de posição, força calculada e estado seguro sem energia.",
  "PWM/CAN e mecanismo de porta.",
  "Carga, ciclos, bloqueio, comando indevido e indicação.")
p("Controles", "Hinges, horns e pushrods", "1 kit", "TBD", "Ferragens aeronáuticas",
  "Transmitir movimentos às superfícies.",
  "Sem folga excessiva, retenção positiva, resistência/corrosão adequadas.",
  "Servos e superfícies.",
  "Carga prova, folga, travamento e inspeção.",
  "Geometria final das superfícies.")

# G — Estrutura, segurança e solo
p("Estrutura", "Longarinas/tubos de fibra de carbono", "1 lote", "TBD crítico", "Tubos/pultrudados com certificado",
  "Caminho principal de carga das asas e booms.",
  "Módulo, resistência, tolerâncias e rastreabilidade; dimensionamento por cargas limite/última.",
  "Junções da asa, fuselagem e naceles.",
  "Cupom, flexão, união e prova estrutural da asa.",
  "Análise estrutural e MTOW fechados.")
p("Estrutura", "Tecido de carbono e epóxi", "1 lote", "TBD", "Sistema laminado aeronáutico",
  "Cascas e reforços estruturais.",
  "Compatibilidade tecido/resina, Tg acima da temperatura de serviço, ficha de cura.",
  "Moldes, núcleo e inserts.",
  "Cupons, teor de resina, cura, porosidade e adesão.")
p("Estrutura", "Núcleo de espuma/honeycomb", "1 lote", "TBD", "PVC estrutural ou honeycomb",
  "Painéis sanduíche leves.",
  "Densidade, compressão/cisalhamento e compatibilidade com resina.",
  "Peles de compósito e inserts.",
  "Cupom sanduíche e inspeção de colagem.")
p("Estrutura", "Compensado aeronáutico/firewall", "1 lote", "TBD crítico", "Compensado birch aeronáutico",
  "Suportar motor traseiro e cargas locais.",
  "Espessura pelo cálculo, colagem certificada, proteção térmica/combustível.",
  "DLE35RA, fuselagem e manta térmica.",
  "Carga de motor, vibração, combustível e temperatura.")
p("Estrutura", "Perfis/chapas de alumínio 7075-T6", "1 lote", "TBD crítico", "Material com certificado",
  "Suportes de motor, tilt e inserts estruturais.",
  "Liga/temperamento certificados; proteção anticorrosiva; raios e espessuras calculados.",
  "Longarinas, mancais, motores e travas.",
  "Inspeção dimensional, certificado e carga prova.")
p("Estrutura", "Fixadores e porcas travantes", "1 lote", "TBD", "Classe aeronáutica, tamanhos por desenho",
  "Uniões desmontáveis.",
  "Resistência, acabamento, controle de torque e retenção positiva.",
  "Todos os subconjuntos.",
  "Rastreabilidade, torque marcado e inspeção 100%.")
p("Estrutura", "Coxins antivibração", "1 kit", "TBD", "Isoladores motor/aviónica",
  "Reduzir transmissão de vibração sem perder rigidez de controle.",
  "Rigidez e frequência natural dimensionadas; compatibilidade com combustível/temperatura.",
  "Motor térmico, sensores e aviónica.",
  "Sweep de vibração, deslocamento máximo e inspeção.")
p("Estrutura", "Trem de pouso/skids", "1 conjunto", "TBD crítico", "Compósito/alumínio absorvente",
  "Apoiar decolagem/pouso VTOL e proteger hélices/fuselagem.",
  "Energia de impacto, estabilidade lateral e altura livre calculadas.",
  "Estrutura central; pontos substituíveis.",
  "Drop test, carga lateral e tombamento.")
p("Estrutura", "Proteção térmica e corta-fogo", "1 kit", "TBD crítico", "Manta/barreira compatível com gasolina",
  "Separar motor/escape/combustível da estrutura e elétrica.",
  "Temperatura, chama e fluidos; bordas seladas.",
  "Firewall, escapamento e compartimento de combustível.",
  "Mapa térmico, inspeção após ensaio de motor e teste de material.")
p("Segurança", "Sistema de paraquedas", "1", "TBD crítico", "Fruity Chutes ou equivalente dimensionado pelo MTOW",
  "Recuperação de emergência.",
  "Capacidade pelo MTOW final, velocidade de descida requerida, método de ejeção e envelope de abertura.",
  "Controlador independente/failsafe; ponto estrutural dedicado.",
  "Extração em solo segura, continuidade, dobra, intervalo de manutenção e carga do ponto.",
  "MTOW, velocidade máxima e análise de risco.", "https://fruitychutes.com/files/Fruity-Chutes-UAV-Product-Guide.pdf")
p("Segurança", "Controlador de terminação/recuperação", "1", "TBD crítico", "Controlador independente",
  "Comandar corte de energia/ignição e recuperação segundo lógica aprovada.",
  "Independência adequada, entradas redundantes, prevenção de acionamento inadvertido.",
  "Kill switch, contatores e paraquedas.",
  "Falhas simples, perda de link, geofence e disparo inadvertido.",
  "Análise de segurança operacional.")
p("Segurança", "Luzes de navegação e anticolisão", "1 kit", "TBD", "LEDs aeronáuticos de alta intensidade",
  "Visibilidade e orientação da aeronave.",
  "Cores/ângulos/intensidade conforme regra aplicável; baixo EMI.",
  "Barramento 12 V/5 V; asas/cauda.",
  "Visibilidade, consumo, temperatura e EMI.")
p("Solo", "Carregador/balanceador 12S", "1", "TBD crítico", "Carregador CC/CV com logging",
  "Carregar o pack fora da aeronave com controle.",
  "12S Li-ion, corrente ajustável, balanceamento/telemetria e proteções.",
  "Conector do pack e chicote de balanceamento.",
  "Calibração, corte, logs e teste com pack instrumentado.")
p("Solo", "Gabinete de armazenamento de baterias", "1", "TBD crítico", "Armário resistente ao fogo ventilado",
  "Armazenamento e carga controlada do pack.",
  "Contenção/ventilação, detecção de temperatura/fumaça e localização segura.",
  "Procedimento de carga e resposta a emergência.",
  "Inspeção do local e simulado de emergência.")
p("Solo", "Bancada de empuxo e célula de carga", "1", "TBD crítico", "Bancada ≥15 kgf por motor",
  "Medir curvas de motores/hélices e consumo.",
  "Célula calibrada, contenção física, medição V/I/RPM/temperatura.",
  "Motores VTOL e motor térmico em bancadas apropriadas.",
  "Calibração, repetibilidade e revisão de segurança antes de cada ensaio.")
p("Solo", "Tacômetro e instrumentos elétricos", "1 kit", "TBD", "Tacômetro óptico + analisador DC",
  "Instrumentar testes e manutenção.",
  "Faixas compatíveis com 12S, 200 A e 12.000 rpm; certificados/calibráveis.",
  "Bancada e aeronave.",
  "Calibração anual e comparação cruzada.")
p("Solo", "Estação de controle em solo", "1", "TBD", "Notebook robusto com Mission Planner/QGroundControl",
  "Planejamento, telemetria, logs e supervisão.",
  "Autonomia, brilho, portas USB, armazenamento e backup.",
  "Rádio RFD USB e rede 4G/Wi-Fi.",
  "Missão simulada, perda de link e restauração de configuração.")


TECH_SPEC = [
    ("Tipo", "UAV cargueiro/de resgate híbrido VTOL de asa fixa"),
    ("Status", "Configuração conceitual consolidada; não liberada para fabricação ou compra integral"),
    ("Envergadura", "4,20 m"),
    ("Comprimento", "2,85 m (envelope atual; requisito inicial 2,6–3,0 m)"),
    ("Área alar", "2,604 m²"),
    ("Corda raiz / ponta", "0,820 m / 0,420 m"),
    ("Afilamento", "0,512"),
    ("Corda aerodinâmica média", "0,6415 m"),
    ("Alongamento geométrico", "AR ≈ 6,77"),
    ("Empenagem", "Cauda em V, ângulo consolidado de 40°"),
    ("Carga útil", "Até 5 kg — meta"),
    ("Alcance", "400 km — meta a validar por ensaio de consumo e reserva"),
    ("Velocidade CFD de referência", "25 m/s, α = 0°"),
    ("Centro de gravidade alvo", "x ≈ 1,05–1,085 m; referência de momentos x = 1,07 m"),
    ("Propulsão VTOL", "4 motores elétricos 12S em naceles inclináveis"),
    ("Pivôs dianteiros", "x = 0,600 m; inclinação de cruzeiro para frente: −90°"),
    ("Pivôs traseiros", "x = 2,000 m; inclinação de cruzeiro para trás: +90°"),
    ("Propulsão de cruzeiro", "Motor traseiro a gasolina em configuração pusher"),
    ("Geração elétrica", "Starter-generator dedicado; motores VTOL não são assumidos como geradores nesta revisão"),
    ("MTOW", "TBD — gate obrigatório antes da compra de motores, ESCs, pack, tilt, estrutura e paraquedas"),
]


MATURITY = {
    "Arquitetura aceita": "Função e arquitetura aprovadas; modelo comercial ainda pode mudar após integração.",
    "Candidato": "Modelo plausível identificado, aguardando dimensionamento/ensaio.",
    "Candidato — compra bloqueada": "Não emitir pedido antes do gate técnico indicado.",
    "TBD": "Requisito funcional definido; fornecedor e part number não selecionados.",
    "TBD crítico": "Item de segurança/desempenho; requer cálculo, desenho ou análise antes da seleção.",
    "Sob medida": "Deve ser projetado/fabricado a partir de desenho liberado.",
    "Opcional / candidato": "Equipamento de missão opcional, sujeito ao perfil operacional e orçamento.",
}


def safe_name(text: str) -> str:
    text = text.lower().replace("ç", "c").replace("ã", "a").replace("õ", "o")
    text = text.replace("á", "a").replace("é", "e").replace("í", "i").replace("ó", "o").replace("ú", "u")
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")[:70]


def write_markdown():
    OUT.mkdir(exist_ok=True)
    DS_DIR.mkdir(exist_ok=True)
    readme = f"""# Tech Spec — UAV híbrido VTOL cargueiro/de resgate

Pacote consolidado em {date.today().strftime('%d/%m/%Y')}.

## Conteúdo

- `Tech_Spec_UAV_Hybrid_VTOL.docx`: documento mestre para revisão.
- `TECH_SPEC.md`: especificação técnica controlada em texto.
- `BOM_Compras.csv`: lista mestra de compras.
- `Data sheets/`: uma ficha técnica para cada uma das {len(PARTS)} linhas da BOM.

## Regra de maturidade

Nenhum item marcado **compra bloqueada**, **TBD crítico**, **TBD** ou **sob medida** está liberado para pedido. “Arquitetura aceita” significa que a função foi aceita, não que a integração esteja certificada.

## Próximos gates

1. Fechar orçamento de massa, MTOW e centro de gravidade.
2. Medir empuxo/corrente/temperatura do conjunto VTOL.
3. Medir empuxo e consumo do motor a gasolina com hélice pusher.
4. Dimensionar pivôs, atuadores, travas e sensores de indexação.
5. Fechar balanço de energia e pack 12S.
6. Executar análise estrutural e prova de carga.
7. Atualizar CFD com camadas de parede e efeitos propulsivos.

O arquivo FreeCAD de referência permanece em `../uav_cargo_manufacturable/CargoUAV_HybridVTOL_Manufacturable.FCStd`.
"""
    (OUT / "README.md").write_text(readme, encoding="utf-8")

    tech_lines = ["# Especificação técnica consolidada", "", "## Configuração"]
    tech_lines += [f"- **{k}:** {v}" for k, v in TECH_SPEC]
    tech_lines += [
        "", "## Conceito operacional",
        "", "A aeronave decola e pousa verticalmente com quatro rotores elétricos. Após atingir altura e condições seguras, o motor traseiro a gasolina é acionado pelo sistema starter-generator e assume o cruzeiro. O starter-generator pode alimentar o barramento e recarregar a bateria dentro dos limites do pack. Os motores VTOL não são considerados geradores nesta revisão, evitando arrasto, controle complexo e riscos de overspeed.",
        "", "## Intertravamento obrigatório dos tilt-rotors",
        "", "A nacele só pode iniciar o movimento quando, para aquele rotor, houver confirmação independente de: comando de torque zero; RPM abaixo do limite; hélice na janela angular indexada; trava da posição atual liberada; trajetória livre; e alimentação do atuador válida. Ao chegar à nova posição, a trava mecânica deve engatar e dois sinais coerentes devem confirmar posição/trava antes de liberar potência ao motor.",
        "", "Sequência de falha: qualquer discordância congela a transição, mantém ou retorna à última posição mecanicamente segura e impede reenergização do rotor afetado. Os rotores dianteiros giram para frente; os traseiros giram para trás. É proibido girar naceles com hélices em rotação livre significativa.",
        "", "## Estado da validação CFD",
        "", "Caso base OpenFOAM 14, 25 m/s e α=0°, malha de 269.773 células. Resultado preliminar médio entre passos 140–150: Cd≈0,06546, Cl≈−0,01247 e Cm≈0,01091. A malha ainda não possui camadas de parede nem representação das hélices; os coeficientes não liberam o projeto.",
        "", "## Gates de liberação de compras",
        "", "- G1 Massa/CG: MTOW, centros de massa e envelopes fechados.",
        "- G2 VTOL: razão empuxo/peso ≥1,5 e margens elétricas/térmicas medidas.",
        "- G3 Cruzeiro: potência, consumo, alcance de 400 km e reserva demonstrados.",
        "- G4 Tilt: cargas, folgas, indexação, travas e análise de falha aprovadas.",
        "- G5 Energia: pack, proteção, redundância e starter-generator validados.",
        "- G6 Estrutura: cargas limite/última, FEA e provas estáticas concluídas.",
        "- G7 Segurança/operação: paraquedas, failsafes, regulamentação e plano de ensaios aprovados.",
        "", "## Rastreabilidade e ressalvas",
        "", "As dimensões acima vêm do modelo paramétrico atual. Cálculos antigos do documento Arduino (como asa de 3 m e massas de 10/19 kg) são históricos e foram superados quando conflitantes. Modelos comerciais listados são candidatos técnicos, não autorização de compra. Confirme revisão, disponibilidade, homologação e ficha oficial diretamente com o fabricante no momento da aquisição.",
    ]
    (OUT / "TECH_SPEC.md").write_text("\n".join(tech_lines) + "\n", encoding="utf-8")

    with (OUT / "BOM_Compras.csv").open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(["ID", "Categoria", "Item", "Quantidade", "Maturidade", "Candidato/modelo", "Arquivo Data sheet"])
        for part in PARTS:
            fn = f"{part.ident}_{safe_name(part.item)}.md"
            writer.writerow([part.ident, part.category, part.item, part.qty, part.maturity, part.candidate, f"Data sheets/{fn}"])

    for part in PARTS:
        fn = DS_DIR / f"{part.ident}_{safe_name(part.item)}.md"
        text = f"""# {part.ident} — {part.item}

## Identificação

| Campo | Valor |
|---|---|
| Categoria | {part.category} |
| Quantidade | {part.qty} |
| Maturidade | {part.maturity} |
| Candidato/modelo | {part.candidate} |
| Liberação de compra | {'BLOQUEADA até fechamento dos gates aplicáveis' if part.maturity != 'Arquitetura aceita' else 'Revisão de integração obrigatória antes do pedido'} |

## Função

{part.purpose}

## Requisitos mínimos de compra

{part.requirements}

## Interfaces

{part.interfaces}

## Critérios de aceitação

{part.acceptance}

## Dependências antes da seleção

{part.dependencies}

## Data sheet de referência

{part.source}

## Documentos que o fornecedor deve entregar

- Data sheet com revisão e part number exatos.
- Desenho dimensional e interfaces.
- Limites elétricos, térmicos, mecânicos e ambientais aplicáveis.
- Curvas de desempenho relevantes e condições do ensaio.
- Certificado de material/conformidade quando aplicável.
- Instruções de instalação, inspeção e vida útil/manutenção.

## Registro de recebimento

- Fabricante / fornecedor: ____________________
- Part number / revisão: ____________________
- Lote / número de série: ____________________
- Data de recebimento: ____________________
- Inspeção de recebimento: APROVADO / REPROVADO
- Responsável: ____________________
"""
        fn.write_text(text, encoding="utf-8")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=8, color=None):
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    pgh = doc.add_heading(text, level=level)
    pgh.paragraph_format.space_before = Pt(10)
    pgh.paragraph_format.space_after = Pt(5)
    return pgh


def add_bullet(doc, text, level=0):
    pgh = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    pgh.add_run(text)
    pgh.paragraph_format.space_after = Pt(2)


def create_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(28)
    styles["Title"].font.color.rgb = RGBColor(20, 54, 82)
    for n in (1, 2, 3):
        styles[f"Heading {n}"].font.name = "Aptos Display"
        styles[f"Heading {n}"].font.color.rgb = RGBColor(20, 54, 82)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("TECH SPEC\nUAV HÍBRIDO VTOL")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Aeronave cargueira e de resgate — especificação e lista mestra de compras")
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(70, 90, 105)
    doc.add_paragraph("")
    box = doc.add_table(rows=4, cols=2)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    box.style = "Table Grid"
    for i, (k, v) in enumerate([
        ("Revisão", "A — configuração conceitual consolidada"),
        ("Data", date.today().strftime("%d/%m/%Y")),
        ("Escopo", "Tech Spec + BOM + fichas técnicas individuais"),
        ("Status", "NÃO LIBERADO PARA FABRICAÇÃO OU COMPRA INTEGRAL"),
    ]):
        set_cell_text(box.cell(i, 0), k, True, 9, (255, 255, 255))
        shade_cell(box.cell(i, 0), "143652")
        set_cell_text(box.cell(i, 1), v, i == 3, 9, (170, 0, 0) if i == 3 else None)
    doc.add_paragraph("")
    warn = doc.add_paragraph()
    warn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = warn.add_run("Documento de engenharia preliminar. Modelos comerciais citados são candidatos e exigem fechamento dos gates técnicos.")
    rr.bold = True
    rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor(150, 80, 0)
    doc.add_page_break()

    add_heading(doc, "1. Finalidade e controle de maturidade", 1)
    doc.add_paragraph("Este documento consolida a configuração atual da aeronave, transforma a arquitetura em requisitos de compra verificáveis e fornece uma ficha técnica individual para cada linha da BOM. A prioridade é impedir compras prematuras de componentes que dependem de MTOW, cargas, energia ou ensaios ainda abertos.")
    for key, desc in MATURITY.items():
        add_bullet(doc, f"{key}: {desc}")

    add_heading(doc, "2. Especificação técnica consolidada", 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_text(table.cell(0, 0), "Parâmetro", True, 9, (255, 255, 255)); shade_cell(table.cell(0,0), "143652")
    set_cell_text(table.cell(0, 1), "Valor / requisito", True, 9, (255, 255, 255)); shade_cell(table.cell(0,1), "143652")
    for idx, (k, v) in enumerate(TECH_SPEC):
        cells = table.add_row().cells
        set_cell_text(cells[0], k, True, 8.5)
        set_cell_text(cells[1], v, False, 8.5)
        if idx % 2:
            shade_cell(cells[0], "EAF0F4"); shade_cell(cells[1], "EAF0F4")

    add_heading(doc, "3. Arquitetura operacional", 1)
    doc.add_paragraph("Decolagem e pouso são verticais com quatro rotores elétricos. Depois de atingir altitude e condições de transição, o motor traseiro a gasolina é acionado e assume o cruzeiro. Um starter-generator dedicado realiza a partida e pode alimentar o barramento/recarregar a bateria dentro dos limites elétricos e térmicos. Nesta revisão, os quatro motores VTOL não são usados como geradores.")
    add_bullet(doc, "Dianteiros: pivôs em x=0,600 m; inclinam para frente até −90° no cruzeiro.")
    add_bullet(doc, "Traseiros: pivôs em x=2,000 m; inclinam para trás até +90° no cruzeiro.")
    add_bullet(doc, "Cruzeiro: motor a gasolina traseiro em configuração pusher.")
    add_bullet(doc, "Carga útil: meta de até 5 kg; alcance: meta de 400 km.")

    add_heading(doc, "4. Intertravamento dos tilt-rotors", 1)
    doc.add_paragraph("O movimento de uma nacele é proibido enquanto a hélice correspondente não estiver parada, indexada e protegida contra reenergização. A lógica deve ser independente por rotor e tolerante a uma falha simples detectável.")
    for line in [
        "Comandar torque zero e confirmar desligamento do ESC.",
        "Confirmar RPM abaixo do limite seguro por sensor independente.",
        "Indexar a hélice dentro da janela angular definida no CAD.",
        "Confirmar caminho livre e liberar a trava mecânica da posição atual.",
        "Mover a nacele com velocidade/torque monitorados.",
        "Engatar a trava mecânica da nova posição e confirmar com dois sinais coerentes.",
        "Somente então permitir nova energização do motor.",
    ]:
        add_bullet(doc, line)
    doc.add_paragraph("Falha ou discordância: interromper a transição, manter/retornar à última posição mecanicamente segura e bloquear a reenergização do rotor afetado. O mecanismo não pode depender apenas do torque de retenção do atuador.")

    add_heading(doc, "5. Orçamento de massa e centro de gravidade", 1)
    doc.add_paragraph("O MTOW ainda não está fechado. O alvo de CG é x≈1,05–1,085 m, com referência de momentos em x=1,07 m. O marcador antigo em x=1,35 m e massas históricas do documento Arduino não são válidos para liberação. A planilha de massa deve conter massa, tolerância e posição x/y/z de cada item, combustível em estados cheio/reserva/vazio e carga útil em todos os envelopes.")
    add_bullet(doc, "Gate G1: nenhum conjunto VTOL, pack, tilt, estrutura ou paraquedas deve ser comprado antes do fechamento de MTOW/CG.")
    add_bullet(doc, "Critério inicial do VTOL: empuxo total medido ≥1,5× o peso máximo de decolagem.")

    add_heading(doc, "6. Energia e geração", 1)
    doc.add_paragraph("A arquitetura conceitual usa barramento 12S. O pack 12S4P com 48 células P45B é apenas candidato: 43,2 V e 18 Ah nominais, cerca de 777,6 Wh e 3,31 kg somente em células. Interconexões, contenção, BMS/monitor, contatores, fusíveis e margem térmica aumentam massa e devem ser projetados por profissional qualificado.")
    add_bullet(doc, "A aviónica essencial possui alimentação de reserva independente e OR-ing.")
    add_bullet(doc, "O Raspberry Pi recebe DC/DC dedicado e redundante; não depende apenas da saída do power module.")
    add_bullet(doc, "A regeneração do starter-generator deve limitar tensão/corrente e suportar rejeição súbita de carga.")

    add_heading(doc, "7. Validação disponível e limitações", 1)
    doc.add_paragraph("O caso CFD base usa OpenFOAM 14 a 25 m/s e α=0°, com 269.773 células. O resultado preliminar médio reportado entre passos 140–150 foi Cd≈0,06546, Cl≈−0,01247 e Cm≈0,01091. Como não há camadas de parede nem representação das hélices/esteiras, os números servem para verificar o fluxo de trabalho, não para congelar o projeto.")

    add_heading(doc, "8. Gates de liberação", 1)
    gates = [
        ("G1 — Massa/CG", "MTOW, centros de massa, combustível e carga útil fechados."),
        ("G2 — VTOL", "Empuxo, corrente, temperatura, vibração e falha de um motor demonstrados."),
        ("G3 — Cruzeiro", "Empuxo, consumo, partida em voo, alcance de 400 km e reserva demonstrados."),
        ("G4 — Tilt", "Pivôs, cargas, folgas, indexação, travas e análise de falhas aprovados."),
        ("G5 — Energia", "Pack, proteções, redundância e starter-generator validados."),
        ("G6 — Estrutura", "FEA, cargas limite/última e provas estáticas concluídas."),
        ("G7 — Segurança", "Paraquedas, failsafes, regulamentação e plano de ensaios aprovados."),
    ]
    for k, v in gates:
        add_bullet(doc, f"{k}: {v}")

    doc.add_section(WD_ORIENT.LANDSCAPE)
    sec2 = doc.sections[-1]
    sec2.orientation = WD_ORIENT.LANDSCAPE
    sec2.page_width, sec2.page_height = sec.page_height, sec.page_width
    sec2.top_margin = Cm(1.2); sec2.bottom_margin = Cm(1.2); sec2.left_margin = Cm(1.2); sec2.right_margin = Cm(1.2)
    add_heading(doc, f"9. Lista mestra de compras — {len(PARTS)} linhas", 1)
    doc.add_paragraph("Cada linha possui ficha técnica individual na pasta “Data sheets”. Quantidade de materiais a granel e ferragens deve ser convertida em quantidade final após desenhos de fabricação.")
    headers = ["ID", "Categoria", "Item", "Qtd.", "Maturidade", "Candidato/modelo", "Liberação"]

    def start_bom_table():
        table = doc.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for col_idx, header in enumerate(headers):
            set_cell_text(table.cell(0, col_idx), header, True, 7.2, (255,255,255))
            shade_cell(table.cell(0, col_idx), "143652")
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat_header = OxmlElement("w:tblHeader")
        repeat_header.set(qn("w:val"), "1")
        tr_pr.append(repeat_header)
        return table

    bom = start_bom_table()
    for idx, part in enumerate(PARTS):
        # Divisão controlada: impede linha partida e repete o cabeçalho na página 2 da BOM.
        if idx == 35:
            doc.add_page_break()
            bom = start_bom_table()
        row = bom.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "1")
        tr_pr.append(cant_split)
        cells = row.cells
        vals = [part.ident, part.category, part.item, part.qty, part.maturity, part.candidate,
                "Revisão" if part.maturity == "Arquitetura aceita" else "Bloqueada"]
        for col_idx, val in enumerate(vals):
            set_cell_text(cells[col_idx], val, col_idx in (0,6), 6.8,
                          (160,0,0) if col_idx == 6 and val == "Bloqueada" else None)
            if idx % 2:
                shade_cell(cells[col_idx], "EAF0F4")
    doc.add_section(WD_ORIENT.PORTRAIT)
    sec3 = doc.sections[-1]
    sec3.orientation = WD_ORIENT.PORTRAIT
    sec3.page_width, sec3.page_height = sec.page_width, sec.page_height
    sec3.top_margin = Cm(1.6); sec3.bottom_margin = Cm(1.6); sec3.left_margin = Cm(1.7); sec3.right_margin = Cm(1.7)
    add_heading(doc, "10. Rastreabilidade e fontes", 1)
    doc.add_paragraph("Fontes principais: modelo FreeCAD atual, caso OpenFOAM atual, conversa de consolidação e documentação histórica do projeto Arduino. Quando houve conflito, prevaleceu a geometria paramétrica atual; números antigos de asa/massa foram tratados como superados. Os links oficiais completos estão nas fichas técnicas individuais.")
    add_heading(doc, "11. Próximas decisões", 1)
    for line in [
        "Preencher orçamento de massa/CG com tolerâncias.",
        "Construir bancada e medir os conjuntos de propulsão.",
        "Dimensionar mecanicamente cada pivô/trava e congelar a janela de indexação das pás.",
        "Fechar o balanço de energia, a estratégia de recarga e o pack.",
        "Atualizar CFD e executar FEA/provas estruturais.",
        "Revisar a BOM após cada gate e somente então liberar part numbers.",
    ]:
        add_bullet(doc, line)

    # Cabeçalho e rodapé independentes por seção para manter alinhamento em
    # páginas retrato e paisagem.
    for section in doc.sections:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        hp = section.header.paragraphs[0]
        hp.text = "TECH SPEC — UAV HÍBRIDO VTOL | REV. A"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.size = Pt(7.5)
        hp.runs[0].font.color.rgb = RGBColor(90, 105, 115)
        fp = section.footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.add_run("Documento preliminar — compra condicionada aos gates técnicos | ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        fp._p.append(fld)
        for run in fp.runs:
            run.font.size = Pt(7.5)

    doc.core_properties.title = "Tech Spec — UAV Híbrido VTOL"
    doc.core_properties.subject = "Especificação técnica e lista mestra de compras"
    doc.core_properties.author = "Projeto Drone — documento consolidado"
    doc.save(OUT / "Tech_Spec_UAV_Hybrid_VTOL.docx")


def make_zip():
    zip_base = ROOT / "Tech_Spec_UAV_Hybrid_VTOL_Package"
    if (zip_base.with_suffix(".zip")).exists():
        (zip_base.with_suffix(".zip")).unlink()
    shutil.make_archive(str(zip_base), "zip", ROOT, "Tech Spec")


if __name__ == "__main__":
    write_markdown()
    create_docx()
    make_zip()
    print(f"Generated {len(PARTS)} datasheets in {OUT}")
